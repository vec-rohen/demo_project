"""Import a TIA Portal PLC tag table (Excel export) as I/O points.

In TIA Portal: PLC tags → open a tag table (or "Show all tags") → Export (the
Excel icon). The sheet has columns such as Name, Path, Data Type, Logical
Address, Comment, Hmi Visible, Hmi Accessible, Hmi Writeable.

Each row becomes an S7 point. The kind is a starting guess from the address
(%I → DI/AI, %Q → DO/AO, %M/DB → DI/AI, or DO/SP when "Hmi Writeable" is
ticked); ranges, units and alarms are added afterwards in the editor.
"""

from __future__ import annotations

import io
import re
from typing import Any

from src.hmi.drivers.s7addr import S7AddressError, check_type, parse

# TIA data type → (point type, S7 type)
TYPE_MAP = {
    "bool": ("bool", None),
    "byte": ("float", "byte"), "usint": ("float", "usint"), "sint": ("float", "sint"),
    "int": ("float", "int"), "uint": ("float", "uint"), "word": ("float", "word"),
    "dint": ("float", "dint"), "udint": ("float", "udint"), "dword": ("float", "dword"),
    "real": ("float", "real"),
}
TRUE = ("true", "yes", "1", "x", "wahr")


class TiaImportError(ValueError):
    pass


def _clean_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", name.strip()).strip("_") or "Tag"
    if base[0].isdigit():
        base = f"T_{base}"
    base = base[:60]
    candidate, n = base, 2
    while candidate in used:
        candidate, n = f"{base}_{n}", n + 1
    used.add(candidate)
    return candidate


def parse_tag_table(data: bytes) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    """Return (points, skipped rows with reasons)."""
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise TiaImportError("Excel import needs the openpyxl package on the portal server") from None
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any parse failure means "not an xlsx"
        raise TiaImportError("That file isn't an Excel workbook (.xlsx) exported from TIA Portal") from None

    header, rows = None, None
    for sheet in wb.worksheets:
        all_rows = list(sheet.iter_rows(values_only=True))
        for i, row in enumerate(all_rows[:10]):
            cells = [str(c).strip().lower() if c is not None else "" for c in row]
            if "name" in cells and "logical address" in cells:
                header, rows = cells, all_rows[i + 1:]
                break
        if header:
            break
    if not header:
        raise TiaImportError("No PLC tag table found – expected columns 'Name', 'Data Type' and 'Logical Address'")

    col = {name: header.index(name) for name in ("name", "data type", "logical address", "comment", "hmi writeable")
           if name in header}
    if "data type" not in col:
        raise TiaImportError("The tag table has no 'Data Type' column")

    def cell(row, key):
        i = col.get(key)
        return "" if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()

    points, skipped, used = [], [], set()
    for row in rows:
        name = cell(row, "name")
        if not name:
            continue
        dtype, address = cell(row, "data type").lower(), cell(row, "logical address")
        if dtype not in TYPE_MAP:
            skipped.append({"name": name, "reason": f"data type {cell(row, 'data type') or '?'} isn't supported on the HMI"})
            continue
        if not address:
            skipped.append({"name": name, "reason": "no logical address (e.g. an optimized DB variable)"})
            continue
        try:
            addr = parse(address)
            point_type, s7_type = TYPE_MAP[dtype]
            if s7_type:
                check_type(addr, s7_type)
            elif addr.size != "X":
                raise S7AddressError(f"Bool at {address} needs a bit address")
        except S7AddressError as exc:
            skipped.append({"name": name, "reason": str(exc)})
            continue
        writable = cell(row, "hmi writeable").lower() in TRUE and addr.area in ("MK", "DB")
        if addr.area == "PE":
            kind = "DI" if point_type == "bool" else "AI"
        elif addr.area == "PA":
            kind = "DO" if point_type == "bool" else "AO"
        elif writable:
            kind = "DO" if point_type == "bool" else "SP"
        else:
            kind = "DI" if point_type == "bool" else "AI"
        point: dict[str, Any] = {
            "name": _clean_name(name, used), "kind": kind, "type": point_type,
            "description": cell(row, "comment") or name,
            "s7": {"address": addr.text(), **({"type": s7_type} if s7_type else {})},
        }
        if point_type == "float":
            point["decimals"] = 1 if s7_type == "real" else 0
            point["s7"]["scale"] = 1
        if writable:
            point["writable"] = True
        points.append(point)
    if not points and not skipped:
        raise TiaImportError("The tag table is empty")
    return points, skipped
